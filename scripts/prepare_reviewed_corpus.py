from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path


INCLUDE_FALSE = {"n", "no", "false", "0", "exclude", "excluded", "제외", "미포함"}


def read_jsonl(path: Path) -> list[dict]:
    with path.open("r", encoding="utf-8") as f:
        return [json.loads(line) for line in f if line.strip()]


def write_jsonl(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="\n") as f:
        for row in rows:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")


def normalize_cell(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def load_review_table(path: Path) -> dict[str, dict]:
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [normalize_cell(v) for v in next(rows)]
        table = {}
        for values in rows:
            row = dict(zip(headers, values))
            chunk_id = normalize_cell(row.get("chunk_id"))
            if chunk_id:
                table[chunk_id] = {k: normalize_cell(v) for k, v in row.items()}
        return table

    with path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        return {
            normalize_cell(row.get("chunk_id")): {k: normalize_cell(v) for k, v in row.items()}
            for row in reader
            if normalize_cell(row.get("chunk_id"))
        }


def include_decision(review_row: dict | None) -> tuple[bool, str]:
    if not review_row:
        return True, "no_review_row_default_include"
    raw = normalize_cell(review_row.get("포함여부"))
    if raw.lower() in INCLUDE_FALSE:
        return False, raw
    return True, raw or "Y"


def main() -> None:
    parser = argparse.ArgumentParser(description="Apply chunk review decisions without rebuilding chunks.")
    parser.add_argument("--source-dir", default="conveyor_rag_parent_child_chunks_v1")
    parser.add_argument("--review-file", default=None)
    parser.add_argument("--output-dir", default="output/rag_reviewed")
    args = parser.parse_args()

    source_dir = Path(args.source_dir)
    review_file = Path(args.review_file) if args.review_file else source_dir / "rag_chunk_review.xlsx"
    output_dir = Path(args.output_dir)

    parents = read_jsonl(source_dir / "parent_chunks.jsonl")
    children = read_jsonl(source_dir / "child_chunks.jsonl")
    reviews = load_review_table(review_file)

    reviewed_children = []
    decisions = []
    for child in children:
        chunk_id = child["chunk_id"]
        review = reviews.get(chunk_id)
        include, raw_decision = include_decision(review)
        decisions.append(
            {
                "chunk_id": chunk_id,
                "parent_id": child["parent_id"],
                "include": "Y" if include else "N",
                "raw_include_value": raw_decision,
                "review_status": normalize_cell(review.get("검수상태")) if review else "",
                "reviewer_note": normalize_cell(review.get("검수자메모")) if review else "",
            }
        )
        if include:
            updated = dict(child)
            if review:
                updated["review_status"] = normalize_cell(review.get("검수상태"))
                updated["reviewer_note"] = normalize_cell(review.get("검수자메모"))
            reviewed_children.append(updated)

    kept_parent_ids = {child["parent_id"] for child in reviewed_children}
    reviewed_parents = [parent for parent in parents if parent["parent_id"] in kept_parent_ids]

    write_jsonl(output_dir / "child_chunks.reviewed.jsonl", reviewed_children)
    write_jsonl(output_dir / "parent_chunks.reviewed.jsonl", reviewed_parents)
    write_jsonl(output_dir / "rag_parent_child_chunks.reviewed.jsonl", reviewed_parents + reviewed_children)

    with (output_dir / "review_decisions.csv").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(decisions[0].keys()))
        writer.writeheader()
        writer.writerows(decisions)

    summary = {
        "source_dir": str(source_dir),
        "review_file": str(review_file),
        "source_parent_count": len(parents),
        "source_child_count": len(children),
        "review_row_count": len(reviews),
        "kept_parent_count": len(reviewed_parents),
        "kept_child_count": len(reviewed_children),
        "excluded_child_count": len(children) - len(reviewed_children),
        "note": "Chunks were not regenerated; only review include/exclude decisions were applied.",
    }
    (output_dir / "reviewed_corpus_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
